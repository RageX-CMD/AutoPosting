#!/usr/bin/env python3
"""Erzeugt aus Spalte C der Event-Excel ein Forum-DOCX.

Das DOCX ist in Copy-Paste-Bloecke aufgeteilt, analog zum Upload-Limit
im deutschen Naruto-Online-Forum (wangEditor). Jeder Block entspricht
einem Forum-Beitrag und liegt zusaetzlich als eigene Datei vor.
"""

from __future__ import annotations

import argparse
import html
import re
import sys
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from openpyxl import load_workbook


# Forum-Posts in der Praxis: ca. 7.600-9.100 Zeichen bzw. 6 Events.
# Am Event-Rand splitten, knapp unter dem Limit das im Forum schon gehalten hat.
CHAR_LIMIT = 9300
MAX_EVENTS_PER_CHUNK = 6
MAX_EVENTS_FIRST_CHUNK = 6
WEEK_SHEET_RE = re.compile(r"^\d{8}$")

DISCORD_URL = "https://discord.gg/jJ439tV2Jp"
DISCORD_BANNER_URL = (
    "https://forum-narutode.narutowebgame.com/api/editor/get-img?"
    "img_name=editor%252F2025-07-23%252F61e982992106b5cb079e9b45eb761d69"
)
MAINTENANCE_START = "10:30 Uhr"

EVENT_HEADER_RE = re.compile(r"^-\s*\d+\s*-")
MAINTENANCE_ITEM_RE = re.compile(r"^\d+\.\s+")
TITLE_RE = re.compile(
    r"Serverwartung\s*&\s*Aktionen\s+(\d{1,2})[.\-/](\d{1,2})[.\-/](\d{2,4})",
    re.IGNORECASE,
)


@dataclass
class EventBlock:
    title: str
    lines: list[str | None] = field(default_factory=list)

    def text(self) -> str:
        parts = [self.title, *(line or "" for line in self.lines)]
        return "\n".join(parts)


@dataclass
class EventPost:
    title: str
    maintenance: list[str]
    events: list[EventBlock]
    source_sheet: str
    source_file: str

    @property
    def date_label(self) -> str:
        match = TITLE_RE.search(self.title or "")
        if not match:
            return datetime.now().strftime("%d.%m.")
        day, month, year = match.groups()
        return f"{int(day):02d}.{int(month):02d}."

    @property
    def iso_date(self) -> str:
        match = TITLE_RE.search(self.title or "")
        if not match:
            if WEEK_SHEET_RE.fullmatch(self.source_sheet):
                return self.source_sheet
            return datetime.now().strftime("%Y%m%d")
        day, month, year = match.groups()
        year_i = int(year)
        if year_i < 100:
            year_i += 2000
        return f"{year_i:04d}{int(month):02d}{int(day):02d}"

    @property
    def folder_date(self) -> str:
        if WEEK_SHEET_RE.fullmatch(self.source_sheet):
            return format_week_label(self.source_sheet)
        return datetime.strptime(self.iso_date, "%Y%m%d").strftime("%d.%m.%Y")


@dataclass
class PostChunk:
    index: int
    total: int
    paragraphs: list[tuple[str, str]]
    event_from: int | None
    event_to: int | None
    char_count: int


def weekly_sheets(workbook) -> list[str]:
    """Jede Woche liegt als eigenes Blatt YYYYMMDD in der Excel."""
    return [name for name in workbook.sheetnames if WEEK_SHEET_RE.fullmatch(name)]


def format_week_label(sheet_name: str) -> str:
    return datetime.strptime(sheet_name, "%Y%m%d").strftime("%d.%m.%Y")


def pick_sheet(workbook, explicit: str | None) -> str:
    weeks = weekly_sheets(workbook)
    if explicit:
        if explicit in workbook.sheetnames:
            return explicit
        if weeks:
            preview = ", ".join(weeks[-6:])
            raise ValueError(f"Wochenblatt {explicit!r} nicht gefunden. Vorhanden u.a.: {preview}")
        raise ValueError(f"Blatt {explicit!r} existiert nicht.")

    if weeks:
        return max(weeks)
    return workbook.sheetnames[-1]


def detect_german_column(worksheet) -> int:
    """Spalte C bevorzugen, sonst die letzte gefuellte Spalte."""
    for col in (3, 2, 1):
        for row in range(1, min(worksheet.max_row, 40) + 1):
            value = worksheet.cell(row, col).value
            if value and str(value).strip():
                return col
    raise ValueError("Keine gefuellte Textspalte gefunden.")


def read_column(worksheet, column: int) -> list[str | None]:
    values: list[str | None] = []
    for row in range(1, worksheet.max_row + 1):
        value = worksheet.cell(row, column).value
        if value is None:
            values.append(None)
        else:
            text = str(value).replace("\xa0", " ").strip()
            values.append(text if text else None)
    while values and values[-1] is None:
        values.pop()
    return values


def normalize_event_title(title: str) -> str:
    match = re.match(r"^-\s*(\d+)\s*-\s*(.*)$", title.strip())
    if not match:
        return title.strip()
    number, name = match.groups()
    return f"- {number} - {name.strip()}"


def parse_post(values: list[str | None], source_file: str, source_sheet: str) -> EventPost:
    nonempty = [(i, text) for i, text in enumerate(values) if text]
    if not nonempty:
        raise ValueError("Spalte ist leer.")

    title_index, title = nonempty[0]
    maintenance: list[str] = []
    events: list[EventBlock] = []
    current: EventBlock | None = None
    seen_event = False

    for text in values[title_index + 1 :]:
        if text is None:
            if current is not None and current.lines and current.lines[-1] is not None:
                current.lines.append(None)
            continue
        if EVENT_HEADER_RE.match(text):
            if current is not None:
                while current.lines and current.lines[-1] is None:
                    current.lines.pop()
            seen_event = True
            current = EventBlock(title=normalize_event_title(text))
            events.append(current)
            continue
        if not seen_event:
            maintenance.append(text)
            continue
        if current is not None:
            current.lines.append(text)

    return EventPost(
        title=title,
        maintenance=maintenance,
        events=events,
        source_sheet=source_sheet,
        source_file=source_file,
    )


def intro_paragraphs(post: EventPost) -> list[tuple[str, str]]:
    date_label = post.date_label
    body = (
        "um euch allen ein besseres Spielerlebnis bieten zu können, werden wir am "
        f"{date_label} eine Serverwartung durchführen. Die Serverwartung findet auf allen "
        f"Servern statt und startet voraussichtlich gegen {MAINTENANCE_START}. Je nach "
        "Umständen könnte die Wartung auch früher oder später erfolgen. Wir werden vor "
        "dem Start auch eine Systemnachricht im Spiel schicken. Während dieser Zeit ist "
        "der Login ins Spiel nicht möglich. Wir bitten daher um eure Nachsicht."
    )
    paragraphs: list[tuple[str, str]] = [
        ("title", post.title),
        ("body", "Liebe Spieler,"),
        ("body", body),
        ("blank", ""),
        ("section", "Offizieller Discord:"),
        ("discord_banner", ""),
        ("blank", ""),
        ("section", "Wartung:"),
    ]
    for item in post.maintenance:
        paragraphs.append(("body", item))
        if not MAINTENANCE_ITEM_RE.match(item):
            continue
        if re.search(r"Fähigkeitsdurchbruch|Themenarena|Neuer Ninja", item, re.IGNORECASE):
            paragraphs.append(("placeholder", "[Wartungs-Bild hier einfügen]"))
    paragraphs.append(("blank", ""))
    paragraphs.append(("section", "Aktionen:"))
    return paragraphs


def event_paragraphs(event: EventBlock) -> list[tuple[str, str]]:
    paragraphs: list[tuple[str, str]] = [
        ("event_title", event.title),
        ("placeholder", "[Event-Bild hier einfügen]"),
        ("blank", ""),
    ]

    def add_blank() -> None:
        if not paragraphs or paragraphs[-1][0] != "blank":
            paragraphs.append(("blank", ""))

    for line in event.lines:
        if line is None:
            add_blank()
            continue

        parts = line.replace("\r\n", "\n").replace("\r", "\n").split("\n")
        for part in parts:
            text = part.strip()
            if not text:
                add_blank()
                continue
            paragraphs.append(("body", text))
            if text.startswith(
                (
                    "Ab Stufe",
                    "Aktionszeitraum:",
                    "Aktionsdetails",
                    "Details zur Aktion",
                )
            ):
                add_blank()

    add_blank()
    return paragraphs


def style_len(paragraphs: list[tuple[str, str]]) -> int:
    return sum(len(text) + 1 for _style, text in paragraphs if text)


def app_base_dir() -> Path:
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


def resource_path(relative: str) -> Path:
    bundle_dir = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent))
    return bundle_dir / relative


def find_excel(explicit: Path | None, search_dir: Path | None = None) -> Path:
    if explicit:
        path = explicit if explicit.is_absolute() else Path.cwd() / explicit
        if not path.exists():
            raise FileNotFoundError(f"Excel nicht gefunden: {path}")
        return path

    folders = []
    if search_dir is not None:
        folders.append(search_dir)
    folders.append(app_base_dir())
    folders.append(Path.cwd())

    seen: set[Path] = set()
    candidates: list[Path] = []
    for folder in folders:
        folder = folder.resolve()
        if folder in seen:
            continue
        seen.add(folder)
        candidates.extend(folder.glob("*.xlsx"))

    candidates = sorted(set(candidates), key=lambda p: p.stat().st_mtime, reverse=True)
    if not candidates:
        raise FileNotFoundError("Keine .xlsx im Programm- oder Projektordner gefunden.")
    return candidates[0]


def list_excel_files(search_dir: Path | None = None) -> list[Path]:
    folder = (search_dir or app_base_dir()).resolve()
    return sorted(folder.glob("*.xlsx"), key=lambda p: p.stat().st_mtime, reverse=True)


def load_weeks(excel: Path) -> list[str]:
    workbook = load_workbook(excel, data_only=True, read_only=True)
    try:
        return weekly_sheets(workbook)
    finally:
        workbook.close()


def week_folder(output_dir: Path, sheet_name: str) -> Path:
    return output_dir / format_week_label(sheet_name)


def week_already_generated(output_dir: Path, sheet_name: str) -> bool:
    folder = week_folder(output_dir, sheet_name)
    return folder.is_dir() and any(folder.glob("*.docx"))


@dataclass
class GenerateResult:
    post: EventPost
    chunks: list[PostChunk]
    combined: Path
    chunk_files: list[Path]
    copy_view: Path
    weeks: list[str]
    week_dir: Path


@dataclass
class BatchGenerateResult:
    created: list[GenerateResult]
    skipped: list[str]
    failed: list[tuple[str, str]]


def _generate_from_workbook(
    workbook,
    excel_name: str,
    output_dir: Path,
    sheet: str | None,
    char_limit: int,
    max_events: int,
) -> GenerateResult:
    weeks = weekly_sheets(workbook)
    sheet_name = pick_sheet(workbook, sheet)
    worksheet = workbook[sheet_name]
    column = detect_german_column(worksheet)
    values = read_column(worksheet, column)

    post = parse_post(values, source_file=excel_name, source_sheet=sheet_name)
    chunks = build_chunks(post, char_limit=char_limit, max_events=max_events)

    week_dir = output_dir / post.folder_date
    week_dir.mkdir(parents=True, exist_ok=True)
    for stale in week_dir.glob("*.docx"):
        stale.unlink()

    combined = week_dir / f"{post.iso_date}_Serverwartung_Aktionen.docx"
    build_combined_document(post, chunks, combined)
    chunk_files = []
    for chunk in chunks:
        path = chunk_filename(week_dir, post, chunk)
        build_single_chunk_document(chunk, path)
        chunk_files.append(path)
    copy_view = week_dir / f"{post.iso_date}_Kopieransicht.html"
    build_copy_view(post, chunks, copy_view)

    return GenerateResult(
        post=post,
        chunks=chunks,
        combined=combined,
        chunk_files=chunk_files,
        copy_view=copy_view,
        weeks=weeks,
        week_dir=week_dir,
    )


def generate_documents(
    excel: Path,
    output_dir: Path,
    sheet: str | None = None,
    char_limit: int = CHAR_LIMIT,
    max_events: int = MAX_EVENTS_PER_CHUNK,
) -> GenerateResult:
    workbook = load_workbook(excel, data_only=True)
    try:
        return _generate_from_workbook(
            workbook,
            excel.name,
            output_dir,
            sheet,
            char_limit,
            max_events,
        )
    finally:
        workbook.close()


def generate_missing_weeks(
    excel: Path,
    output_dir: Path,
    char_limit: int = CHAR_LIMIT,
    max_events: int = MAX_EVENTS_PER_CHUNK,
    progress=None,
) -> BatchGenerateResult:
    workbook = load_workbook(excel, data_only=True)
    created: list[GenerateResult] = []
    skipped: list[str] = []
    failed: list[tuple[str, str]] = []
    try:
        weeks = weekly_sheets(workbook)
        total = len(weeks)
        for index, sheet_name in enumerate(weeks, start=1):
            label = format_week_label(sheet_name)
            prefix = f"[{index}/{total}] {label}"
            if week_already_generated(output_dir, sheet_name):
                skipped.append(sheet_name)
                if progress:
                    progress(f"{prefix}: schon vorhanden, übersprungen")
                continue
            try:
                result = _generate_from_workbook(
                    workbook,
                    excel.name,
                    output_dir,
                    sheet_name,
                    char_limit,
                    max_events,
                )
                created.append(result)
                if progress:
                    progress(f"{prefix}: erstellt ({len(result.chunks)} Beiträge)")
            except Exception as exc:  # noqa: BLE001 - einzelne Woche darf den Lauf nicht stoppen
                failed.append((sheet_name, str(exc)))
                if progress:
                    progress(f"{prefix}: Fehler — {exc}")
    finally:
        workbook.close()

    return BatchGenerateResult(created=created, skipped=skipped, failed=failed)


def build_chunks(
    post: EventPost,
    char_limit: int = CHAR_LIMIT,
    max_events: int = MAX_EVENTS_PER_CHUNK,
) -> list[PostChunk]:
    intro = intro_paragraphs(post)
    event_blocks = [event_paragraphs(event) for event in post.events]

    chunks: list[list[tuple[str, str]]] = []
    ranges: list[tuple[int | None, int | None]] = []
    current = list(intro)
    start_idx = 1 if event_blocks else None
    events_in_chunk = 0

    for index, block in enumerate(event_blocks, start=1):
        prospective = current + block
        over_chars = style_len(prospective) > char_limit
        over_events = events_in_chunk >= max_events
        # Nie mitten im Event splitten, und Intro nicht allein lassen.
        must_split = events_in_chunk > 0 and (over_chars or over_events)

        if must_split:
            chunks.append(current)
            ranges.append((start_idx, index - 1))
            current = list(block)
            start_idx = index
            events_in_chunk = 1
            continue

        current = prospective
        events_in_chunk += 1

    if current:
        chunks.append(current)
        end_idx = len(event_blocks) if event_blocks else None
        ranges.append((start_idx, end_idx))

    result: list[PostChunk] = []
    total = len(chunks)
    for i, paragraphs in enumerate(chunks, start=1):
        event_from, event_to = ranges[i - 1]
        result.append(
            PostChunk(
                index=i,
                total=total,
                paragraphs=paragraphs,
                event_from=event_from,
                event_to=event_to,
                char_count=style_len(paragraphs),
            )
        )
    return result


def set_run_font(run, *, size: int, bold: bool, color: RGBColor | None = None, italic: bool = False) -> None:
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    if color is not None:
        run.font.color.rgb = color


def add_styled_paragraph(document: Document, style: str, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.space_before = Pt(0)

    if style == "blank":
        # Ein wirklich leerer Word-Absatz wird vom Forum-Editor beim ersten
        # Einfügen entfernt. Das geschützte Leerzeichen wird als &nbsp;
        # übernommen und hält damit die gewünschte sichtbare Leerzeile.
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run("\u00a0")
        set_run_font(run, size=11, bold=False)
        return

    if style == "discord_banner":
        image_path = resource_path("assets/discord_banner.png")
        if image_path.exists():
            run = paragraph.add_run()
            run.add_picture(str(image_path), width=Inches(6.2))
            relationship = paragraph.part.relate_to(
                DISCORD_URL,
                RELATIONSHIP_TYPE.HYPERLINK,
                is_external=True,
            )
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(qn("r:id"), relationship)
            paragraph._p.remove(run._r)
            hyperlink.append(run._r)
            paragraph._p.append(hyperlink)
        else:
            run = paragraph.add_run(f"Offizieller Discord: {DISCORD_URL}")
            set_run_font(run, size=11, bold=True)
        paragraph.paragraph_format.space_after = Pt(8)
        return

    if style == "marker":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(28)
        paragraph.paragraph_format.space_after = Pt(6)
        run = paragraph.add_run(text)
        set_run_font(run, size=12, bold=True, color=RGBColor(0xC0, 0x39, 0x2B))
        return

    if style == "marker_sub":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(28)
        run = paragraph.add_run(text)
        set_run_font(run, size=10, bold=False, italic=True, color=RGBColor(0x7F, 0x8C, 0x8D))
        return

    run = paragraph.add_run(text)
    if style == "title":
        set_run_font(run, size=18, bold=True)
        paragraph.paragraph_format.space_after = Pt(12)
    elif style == "section":
        set_run_font(run, size=14, bold=True)
        paragraph.paragraph_format.space_before = Pt(8)
        paragraph.paragraph_format.space_after = Pt(8)
    elif style == "event_title":
        set_run_font(run, size=13, bold=True)
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.space_after = Pt(6)
    elif style == "placeholder":
        set_run_font(run, size=11, bold=False, italic=True, color=RGBColor(0x8E, 0x44, 0xAD))
        paragraph.paragraph_format.space_after = Pt(8)
    else:
        set_run_font(run, size=11, bold=True)


def add_limit_gap(document: Document, chunk: PostChunk, next_chunk: PostChunk) -> None:
    """Sichtbarer Abstand genau dort, wo das Forum-Limit greift."""
    event_range = ""
    if chunk.event_from and chunk.event_to:
        event_range = f"Events {chunk.event_from}–{chunk.event_to}  ·  "
    add_styled_paragraph(
        document,
        "marker",
        "▼  COPY-PASTE GRENZE  ·  FORUM-UPLOAD-LIMIT  ▼",
    )
    add_styled_paragraph(
        document,
        "marker_sub",
        (
            f"Ende Beitrag {chunk.index}/{chunk.total}  ·  {event_range}"
            f"{chunk.char_count} Zeichen  ·  Nächsten Block separat einfügen "
            f"(Beitrag {next_chunk.index}: Events {next_chunk.event_from}–{next_chunk.event_to})"
        ),
    )
    for _ in range(8):
        document.add_paragraph()
    paragraph = document.add_paragraph()
    run = paragraph.add_run()
    run.add_break(WD_BREAK.PAGE)


def write_chunk_body(document: Document, chunk: PostChunk) -> None:
    for style, text in chunk.paragraphs:
        add_styled_paragraph(document, style, text)


def chunk_filename(output_dir: Path, post: EventPost, chunk: PostChunk) -> Path:
    return output_dir / f"{post.iso_date}_forum-post_{chunk.index:02d}.docx"


def build_combined_document(post: EventPost, chunks: list[PostChunk], path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Pt(48)
    section.bottom_margin = Pt(48)
    section.left_margin = Pt(54)
    section.right_margin = Pt(54)

    for i, chunk in enumerate(chunks):
        write_chunk_body(document, chunk)
        if i < len(chunks) - 1:
            add_limit_gap(document, chunk, chunks[i + 1])

    document.save(path)


def build_single_chunk_document(chunk: PostChunk, path: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Pt(48)
    section.bottom_margin = Pt(48)
    section.left_margin = Pt(54)
    section.right_margin = Pt(54)
    write_chunk_body(document, chunk)
    document.save(path)


def forum_html_paragraph(style: str, text: str) -> str:
    escaped = html.escape(text, quote=False)
    if style == "blank":
        # Der alte WangEditor des Forums entfernt Format-Tags innerhalb von
        # <p>. Ein zusätzliches <br> bleibt hingegen als echte Leerzeile.
        return "<br>"
    if style == "title":
        return f"<h2>{escaped}</h2>"
    if style == "section":
        return f"<h3>{escaped}</h3>"
    if style == "event_title":
        return f"<h4>{escaped}</h4>"
    if style == "placeholder":
        return f"<i>{escaped}</i><br>"
    if style == "discord_banner":
        return (
            f'<a href="{html.escape(DISCORD_URL, quote=True)}" target="_blank">'
            f'<img src="{html.escape(DISCORD_BANNER_URL, quote=True)}" '
            'alt="Offizieller Discord"></a><br>'
        )
    # Wichtig: <b> muss auf oberster Ebene stehen. Bei <p><b>...</b></p>
    # löscht der Paste-Filter des Forums das <b> absichtlich.
    return f"<b>{escaped}</b><br>"


def chunk_to_forum_html(chunk: PostChunk) -> str:
    return "\n".join(
        forum_html_paragraph(style, text)
        for style, text in chunk.paragraphs
    )


def build_copy_view(post: EventPost, chunks: list[PostChunk], path: Path) -> None:
    cards: list[str] = []
    for chunk in chunks:
        body = chunk_to_forum_html(chunk)
        event_range = ""
        if chunk.event_from and chunk.event_to:
            event_range = f" · Events {chunk.event_from}–{chunk.event_to}"
        cards.append(
            f"""
            <section class="card">
              <div class="card-head">
                <div>
                  <h2>Beitrag {chunk.index} von {chunk.total}</h2>
                  <span>{chunk.char_count} Zeichen{event_range}</span>
                </div>
                <button type="button" onclick="copyPost('post-{chunk.index}', this)">
                  Beitrag {chunk.index} formatiert kopieren
                </button>
              </div>
              <div class="post" id="post-{chunk.index}">{body}</div>
            </section>
            """
        )

    page_title = html.escape(post.title)
    document = f"""<!doctype html>
<html lang="de">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width,initial-scale=1">
  <title>{page_title} – Kopieransicht</title>
  <style>
    * {{ box-sizing: border-box; }}
    body {{ margin: 0; background: #f4f1ec; color: #2c3e50;
            font-family: "Segoe UI", Arial, sans-serif; }}
    header {{ position: sticky; top: 0; z-index: 10; padding: 18px 24px;
              background: #1e2a39; color: white; box-shadow: 0 2px 10px #0003; }}
    header h1 {{ margin: 0 0 4px; font-size: 22px; }}
    header p {{ margin: 0; color: #d7dde5; }}
    main {{ width: min(960px, calc(100% - 32px)); margin: 24px auto 60px; }}
    .card {{ margin: 0 0 24px; overflow: hidden; background: white;
             border: 1px solid #ddd4ca; border-radius: 8px; }}
    .card-head {{ display: flex; align-items: center; justify-content: space-between;
                  gap: 16px; padding: 14px 18px; background: #f8f6f3;
                  border-bottom: 1px solid #e4ddd5; }}
    .card-head h2 {{ margin: 0 0 3px; font-size: 18px; }}
    .card-head span {{ color: #687786; font-size: 13px; }}
    button {{ border: 0; border-radius: 5px; padding: 10px 15px;
              background: #d35400; color: white; font-weight: 700;
              cursor: pointer; }}
    button:hover {{ background: #e67e22; }}
    button.done {{ background: #218c53; }}
    .post {{ padding: 22px 28px; font-family: Arial, sans-serif;
             color: #111; background: white; }}
    .post p {{ margin: 0 0 7px; line-height: 1.35; }}
  </style>
</head>
<body>
  <header>
    <h1>{page_title}</h1>
    <p>Je Beitrag auf „formatiert kopieren“ klicken und danach direkt ins Forum einfügen.</p>
  </header>
  <main>{''.join(cards)}</main>
  <script>
    async function copyPost(id, button) {{
      const node = document.getElementById(id);
      let ok = false;
      try {{
        if (navigator.clipboard && window.ClipboardItem) {{
          const item = new ClipboardItem({{
            'text/html': new Blob([node.innerHTML], {{type: 'text/html'}}),
            'text/plain': new Blob([node.innerText], {{type: 'text/plain'}})
          }});
          await navigator.clipboard.write([item]);
          ok = true;
        }}
      }} catch (error) {{
        ok = false;
      }}
      if (!ok) {{
        const selection = window.getSelection();
        const range = document.createRange();
        range.selectNodeContents(node);
        selection.removeAllRanges();
        selection.addRange(range);
        ok = document.execCommand('copy');
        selection.removeAllRanges();
      }}
      if (ok) {{
        const old = button.textContent;
        button.textContent = 'Kopiert ✓';
        button.classList.add('done');
        setTimeout(() => {{
          button.textContent = old;
          button.classList.remove('done');
        }}, 1800);
      }} else {{
        alert('Kopieren wurde vom Browser blockiert. Inhalt markieren und Strg+C drücken.');
      }}
    }}
  </script>
</body>
</html>
"""
    path.write_text(document, encoding="utf-8")


def print_summary(
    post: EventPost,
    chunks: list[PostChunk],
    combined: Path,
    output_dir: Path,
    weeks: list[str],
) -> None:
    print(f"Excel:  {post.source_file}")
    if weeks:
        week_note = "neueste Woche" if post.source_sheet == max(weeks) else "manuell gewählt"
        print(f"Woche:  {post.source_sheet} ({format_week_label(post.source_sheet)}, {week_note})")
        print(f"Tabs:   {len(weeks)} Wochenblätter (YYYYMMDD), Spalte C")
    else:
        print(f"Blatt:  {post.source_sheet}")
    print(f"Titel:  {post.title}")
    print(f"Events: {len(post.events)}")
    print(f"Ordner: {output_dir}")
    print(f"Limit:  {CHAR_LIMIT} Zeichen / max. {MAX_EVENTS_PER_CHUNK} Events je Beitrag")
    print()
    for chunk in chunks:
        extra = ""
        if chunk.event_from and chunk.event_to:
            extra = f", Events {chunk.event_from}–{chunk.event_to}"
        print(
            f"  Beitrag {chunk.index}/{chunk.total}: {chunk.char_count} Zeichen{extra} -> "
            f"{chunk_filename(output_dir, post, chunk).name}"
        )
    print()
    print(f"Gesamt-DOCX: {combined}")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Erzeugt ein Forum-Event-DOCX aus Spalte C der Excel-Datei."
    )
    parser.add_argument("--excel", type=Path, help="Pfad zur .xlsx (Standard: neueste im Ordner)")
    parser.add_argument(
        "--sheet",
        help="Wochenblatt YYYYMMDD (Standard: neueste Woche, z.B. 20260917)",
    )
    parser.add_argument(
        "--list-weeks",
        action="store_true",
        help="Vorhandene Wochenblätter auflisten und beenden",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path("output"),
        help="Zielordner (Standard: ./output)",
    )
    parser.add_argument("--char-limit", type=int, default=CHAR_LIMIT)
    parser.add_argument("--max-events", type=int, default=MAX_EVENTS_PER_CHUNK)
    parser.add_argument(
        "--all-missing",
        action="store_true",
        help="Alle Wochen erzeugen, vorhandene Wochenordner überspringen",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    excel_path = find_excel(args.excel)
    if args.list_weeks:
        weeks = load_weeks(excel_path)
        print(f"Excel: {excel_path.name}")
        if not weeks:
            print("Keine Wochenblätter im Format YYYYMMDD gefunden.")
            return 1
        newest = max(weeks)
        for name in weeks:
            marker = "  <- neueste" if name == newest else ""
            print(f"  {name}  {format_week_label(name)}{marker}")
        return 0

    output_dir = args.output_dir if args.output_dir.is_absolute() else app_base_dir() / args.output_dir
    if args.all_missing:
        batch = generate_missing_weeks(
            excel=excel_path,
            output_dir=output_dir,
            char_limit=args.char_limit,
            max_events=args.max_events,
            progress=print,
        )
        print(
            f"Fertig: {len(batch.created)} neu, "
            f"{len(batch.skipped)} übersprungen, "
            f"{len(batch.failed)} Fehler"
        )
        return 1 if batch.failed and not batch.created else 0

    result = generate_documents(
        excel=excel_path,
        output_dir=output_dir,
        sheet=args.sheet,
        char_limit=args.char_limit,
        max_events=args.max_events,
    )
    print_summary(result.post, result.chunks, result.combined, result.week_dir, result.weeks)
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except Exception as exc:  # noqa: BLE001 - CLI soll den Fehler klar anzeigen
        print(f"Fehler: {exc}", file=sys.stderr)
        sys.exit(1)
